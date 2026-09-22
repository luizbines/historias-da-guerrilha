"""Gera as páginas perfs/<slug>-perfil.html a partir de text_victims.xlsx.

Cada quadrado da bandeira (data-victim em bandeira.html) abre o card
overlays/overlay-<slug>.html, cuja seta aponta para perfs/<slug>-perfil.html.
Este script casa cada linha da planilha com um card pelo nome, gera a
página de perfil e aponta a seta do card (visível) para ela.

Arnaldo e Soledad têm páginas próprias e não são gerados aqui.

As traduções (inglês e espanhol) vêm de scripts/traducoes_perfis.json, criado
por scripts/traduzir_perfis.py. Texto sem tradução aparece em português.

Uso (na raiz do repositório):  python3 scripts/gerar_perfis.py
Requer openpyxl.
"""
import difflib
import hashlib
import html
import json
import os
import re
import unicodedata

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = "perfs"
EXCLUDE = {"arnaldo", "soledad"}
# Pessoas com página própria ou sem quadrado na bandeira (linhas sem texto).
SKIP_ROWS = {
    "Arnaldo Cardoso Rocha",
    "Soledad Barret Viedma",
    "Manuel Alves de Oliveira",
    "Gerosina Silva Pereira",
    "Antônio Carlos Silveira Alves",
    "Ângelo Pezzuti da Silva",
}
PARTICLES = {"da", "de", "do", "das", "dos", "e"}
ARROW = '<img src="images/arrow-2-8858.svg" class="arrow-2-8858" alt="arrow-2" />'
TRANSLATIONS = os.path.join(ROOT, "scripts", "traducoes_perfis.json")
LANGS = (("pt", "pt-BR"), ("en", "en"), ("es", "es"))
missing_translations = set()


def text_key(text):
    return hashlib.sha1(text.encode("utf-8")).hexdigest()[:16]


def load_translations():
    if not os.path.exists(TRANSLATIONS):
        return {}
    with open(TRANSLATIONS, encoding="utf-8") as f:
        return json.load(f)


translations = load_translations()


def translate(text, lang):
    if lang == "pt":
        return text
    result = translations.get(lang, {}).get(text_key(text))
    if result is None:
        missing_translations.add((lang, text_key(text)))
        return text
    return result


def norm(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def read_card(slug):
    text = open(os.path.join(ROOT, "overlays", f"overlay-{slug}.html"), encoding="utf-8").read()
    fields = dict(re.findall(r'class="text-(name|location|date|occupation)"><span[^>]*>(.*?)</span>', text))
    fields["img"] = re.search(r'<img[^>]*src="([^"]*)"', text).group(1)
    return {k: html.unescape(v).strip() for k, v in fields.items()}


def flag_slugs():
    text = open(os.path.join(ROOT, "bandeira.html"), encoding="utf-8").read()
    return set(re.findall(r'data-victim="([^"]+)"', text)) - EXCLUDE


def match_rows(cards):
    ws = openpyxl.load_workbook(os.path.join(ROOT, "text_victims.xlsx")).active
    header = [c.value for c in ws[1]]
    assert header == ["nome", "circunstancias da morte", "conclusao da cnv", "biografia"], header
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[0] not in SKIP_ROWS]
    free = set(cards)
    mapping = {}
    # 1ª passada: nomes idênticos; 2ª: grafias diferentes ("Barcelos" x "Barcellos")
    pending = []
    by_name = {norm(cards[s]["name"]): s for s in cards}
    for row in rows:
        # Remove apelidos: "(Bacuri)", "“Alfaiate”"
        key = norm(re.sub(r"\(.*?\)|“.*?”", "", row[0]))
        if by_name.get(key) in free:
            mapping[by_name[key]] = row
            free.discard(by_name[key])
        else:
            pending.append((key, row))
    for key, row in pending:
        score, slug = max((difflib.SequenceMatcher(None, key, norm(cards[s]["name"])).ratio(), s) for s in free)
        if score < 0.7:
            raise SystemExit(f"Sem card para {row[0]!r} (melhor: {slug}, {score:.2f})")
        mapping[slug] = row
        free.discard(slug)
    return mapping, free


def split_paragraphs(text):
    text = (text or "").strip()
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


def paragraphs(text):
    """Um bloco por idioma; perfil-vitima.css mostra só o do idioma atual."""
    blocks = []
    for lang, code in LANGS:
        ps = "\n".join(
            f"            <p>{html.escape(translate(p, lang)).replace(chr(10), '<br>')}</p>"
            for p in split_paragraphs(text)
        )
        blocks.append(f'          <div data-l="{lang}" lang="{code}">\n{ps}\n          </div>')
    return "\n".join(blocks)


def inline_langs(text):
    return "".join(
        f'<span data-l="{lang}" lang="{code}">{html.escape(translate(text, lang))}</span>'
        for lang, code in LANGS
    )


def title_lines(name):
    lines, pending = [], []
    for word in name.split():
        pending.append(word)
        if word.lower() not in PARTICLES:
            lines.append(" ".join(pending))
            pending = []
    if pending:
        lines.append(" ".join(pending))
    return lines


def clean_meta(card):
    location = ", ".join(p for p in card.get("location", "").split(", ") if p and p != "NA")
    date = " - ".join("?" if p == "NA" else p for p in card.get("date", "").split(" - "))
    occupation = card.get("occupation", "")
    if occupation in ("NA", "Nao Consta"):
        occupation = ""
    return location, date, occupation


def section(key, label, body, label_class="highlight-bg"):
    return f"""      <section class="text-section">
        <h2 class="section-label"><span class="{label_class}" data-i18n="{key}">{label}</span></h2>
        <div class="section-body">
{body}
        </div>
      </section>"""


def render(card, row):
    name = card["name"]
    lines = title_lines(name)
    location, date, occupation = clean_meta(card)
    meta = [html.escape(date)] if date else []
    if location:
        meta.append(f'<span class="highlight-bold">{html.escape(location)}</span>')
    if occupation:
        meta.append(inline_langs(occupation))

    sections = []
    if row[3] and row[3].strip():
        sections.append(section("profile.biography", "BIOGRAFIA", paragraphs(row[3])))
    sections.append(section("profile.circumstances", "CIRCUNSTÂNCIAS DA MORTE", paragraphs(row[1])))
    conclusion = section("profile.cnv_conclusion", "CONCLUSÃO DA CNV", paragraphs(row[2]), "highlight-black-bg")

    return f"""<!DOCTYPE html>
<!-- Gerado por scripts/gerar_perfis.py a partir de text_victims.xlsx. Edite a planilha e rode o script. -->
<html lang="pt-br">

<head>
  <meta charset="utf-8" />
  <meta content="width=device-width, initial-scale=1.0" name="viewport" />
  <title>{html.escape(name)} - Perfil</title>
  <script src="../i18n.js?v=2"></script>
  <link
    href="https://fonts.googleapis.com/css2?family=Roboto+Mono:ital,wght@0,300;0,400;0,700;1,400&amp;family=Playfair+Display:ital,wght@0,600;1,400&amp;display=swap"
    rel="stylesheet" />
  <link href="perfil-vitima.css?v=1" rel="stylesheet" />
</head>

<body>
  <header class="global-header">
    <a href="../index.html" data-i18n="nav.home">INÍCIO</a>
    <a href="../projeto.html" data-i18n="nav.project">O PROJETO</a>
    <a href="../comiss-o.html" data-i18n="nav.commission">A COMISSÃO</a>
    <a href="../quem-pagina.html" data-i18n="nav.others">OS OUTROS</a>
  </header>

  <main class="profile-container">

    <div class="hero-section">
      <div class="hero-top">
        <div class="hero-photo">
          <img class="hero-image" alt="Foto de {html.escape(name)}" src="../{html.escape(card['img'])}" />
        </div>
        <div class="hero-header-right">
          <h1 class="hero-title" style="--title-chars: {max(len(l) for l in lines)};">{'<br>'.join(html.escape(l) for l in lines)}</h1>
          <div class="hero-meta-info">
            {'<br>'.join(meta)}
          </div>
        </div>
      </div>
    </div>

{chr(10).join(sections)}

    <div class="red-banner-section">
      <div class="red-banner-container">
{conclusion}
      </div>
    </div>

  </main>

  <iframe src="../bandeira.html?v=17" scrolling="no"
    style="width: 100%; height: 100vh; border: none; margin-top: 80px; display: block; overflow: hidden;"></iframe>

  <div style="background-color: #f5f2e6; text-align: center; overflow: hidden; padding-top: 50px;">
    <h1
      style="margin: 0; margin-bottom: -3vw; font-size: 7vw; font-weight: 700; color: #000000; letter-spacing: 1.5vw; padding-left: 1.5vw; white-space: nowrap; padding-bottom: 0; border-bottom: none;" data-i18n="index.motto">
      ORDEM E PROGRESSO
    </h1>
  </div>

  <footer class="credits"
    style="padding: 60px 5vw; text-align: left; background-color: #000000; color: #ffffff; font-family: 'Roboto Mono', monospace;">
    <div style="margin-bottom: 40px;">
      <span style="font-size: 28px; font-weight: bold; letter-spacing: 1px; display: block; margin-bottom: 10px;" data-i18n="footer.title">
        HISTÓRIAS DA GUERRILHA
      </span>
      <span style="font-size: 18px; font-weight: bold; letter-spacing: 1px; display: block;" data-i18n="footer.subtitle">
        VIDAS PERDIDAS NA DITADURA CIVIL MILITAR BRASILEIRA
      </span>
    </div>
    <div style="font-size: 14px; line-height: 2.2;">
      <span data-i18n="footer.developed">Projeto desenvolvido por Maria Eduarda Rocha</span><br>
      <span data-i18n="footer.orientation">Orientação Barbara Castro</span><br>
      <span data-i18n="footer.institution">Escola Superior de Desenho Industrial - ESDI, UERJ</span>
    </div>
  </footer>
</body>

</html>
"""


def link_card(slug):
    path = os.path.join(ROOT, "overlays", f"overlay-{slug}.html")
    text = open(path, encoding="utf-8").read()
    new = re.sub(r'(<a class="arrow-link" href=")[^"]*(")', rf"\g<1>{OUT_DIR}/{slug}-perfil.html\2", text)
    new = re.sub(r'(<a class="arrow-link"[^>]*>)\s*(</a>)', rf"\1\n{ARROW}\n\2", new)
    if new != text:
        open(path, "w", encoding="utf-8").write(new)


def main():
    cards = {slug: read_card(slug) for slug in flag_slugs()}
    mapping, without_text = match_rows(cards)
    os.makedirs(os.path.join(ROOT, OUT_DIR), exist_ok=True)
    for slug, row in sorted(mapping.items()):
        with open(os.path.join(ROOT, OUT_DIR, f"{slug}-perfil.html"), "w", encoding="utf-8") as f:
            f.write(render(cards[slug], row))
        link_card(slug)
    print(f"{len(mapping)} páginas geradas.")
    if missing_translations:
        print(f"{len(missing_translations)} textos sem tradução (rode scripts/traduzir_perfis.py).")
    if without_text:
        print(f"{len(without_text)} quadrados da bandeira sem texto na planilha:")
        for slug in sorted(without_text):
            print(f"  {slug}  ({cards[slug]['name']})")


if __name__ == "__main__":
    main()
